"""Turn whatever a vendor sent into model-readable content blocks, keeping location anchors
(cell refs, paragraph numbers, page images) so every extracted number can point back to its source."""
import base64
from pathlib import Path

from openpyxl import load_workbook
from docx import Document

IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp", ".gif": "image/gif"}


def _b64(data: bytes) -> str:
    return base64.standard_b64encode(data).decode("utf-8")


def xlsx_to_text(data: bytes) -> str:
    import io
    wb = load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== Sheet: {ws.title} ===")
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}: {c.value}" for c in row if c.value not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def docx_to_text(data: bytes) -> str:
    import io
    doc = Document(io.BytesIO(data))
    out = [f"[para {i}] {p.text}" for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
    for ti, t in enumerate(doc.tables, 1):
        for ri, row in enumerate(t.rows, 1):
            out.append(f"[table {ti} row {ri}] " + " | ".join(c.text for c in row.cells))
    return "\n".join(out)


def to_blocks(name: str, data: bytes):
    """Returns a list of content blocks for one file, each preceded by a label naming the file."""
    ext = Path(name).suffix.lower()
    label = {"type": "text", "text": f"--- FILE: {name} ---"}
    if ext == ".pdf":
        return [label, {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": _b64(data)}}]
    if ext in IMAGE_TYPES:
        return [label, {"type": "image", "source": {"type": "base64", "media_type": IMAGE_TYPES[ext], "data": _b64(data)}}]
    if ext in (".xlsx", ".xlsm"):
        return [label, {"type": "text", "text": "Spreadsheet contents, each cell prefixed with its reference:\n" + xlsx_to_text(data)}]
    if ext == ".docx":
        return [label, {"type": "text", "text": "Word document, paragraphs numbered:\n" + docx_to_text(data)}]
    # .txt / .eml / anything else: treat as text
    return [label, {"type": "text", "text": data.decode("utf-8", errors="replace")}]
