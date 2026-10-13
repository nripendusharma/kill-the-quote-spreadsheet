"""
Builds the fabricated demo dataset for "Kill the Quote Spreadsheet".

Outputs (all fictional companies and people):
  rfx.json                    the buyer's RFx: 30 corrugated lines, questionnaire, terms, assumptions
  last_year_contract.csv      incumbent's 2025 contract prices (needed for "rest same as last year")
  vendor_responses/           five replies, each messy in its own way
  attachments/                certificates and test reports vendors attached
  answer_key.json             ground truth for the extraction scorecard (never shown to the AI)

Run: python build_dataset.py   (deterministic, seeded)
"""
import json, math, random, csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageDraw, ImageFont, ImageFilter

random.seed(14)
ROOT = Path(__file__).parent
VR = ROOT / "vendor_responses"; VR.mkdir(exist_ok=True)
AT = ROOT / "attachments"; AT.mkdir(exist_ok=True)
pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/arial.ttf"))
pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/arialbd.ttf"))

# ---------------------------------------------------------------- board maths
TAKE_UP = {"B": 1.35, "C": 1.45}
GRADES = {  # layer GSM: liners and flutes alternate, starting with a liner
    3: [150, 120, 150],
    5: [180, 120, 150, 120, 180],
    7: [200, 140, 150, 140, 150, 140, 200],
}
FLUTES = {3: None, 5: "BC", 7: "BCB"}


def board_gsm(ply, flute):
    layers = GRADES[ply]
    fl = list(flute) if len(flute) > 1 else [flute] * (ply // 2)
    total, fi = 0, 0
    for i, g in enumerate(layers):
        if i % 2 == 0:
            total += g
        else:
            total += g * TAKE_UP[fl[fi]]; fi += 1
    return round(total)


def rsc_blank_m2(L, W, H):
    return (2 * L + 2 * W + 35) * (W + H) / 1e6


# id, description, kind, ply, flute, L, W, H, area_override_m2, base_qty
LINES = [
    ("L01", "Biscuit shipper", "RSC", 3, "B", 380, 250, 200, None, 600000),
    ("L02", "Snack shipper", "RSC", 3, "C", 450, 300, 250, None, 450000),
    ("L03", "Soap outer carton", "RSC", 3, "B", 320, 220, 180, None, 500000),
    ("L04", "Instant noodle shipper", "RSC", 3, "B", 400, 280, 220, None, 550000),
    ("L05", "Tea pack outer", "RSC", 3, "C", 350, 250, 250, None, 300000),
    ("L06", "Spice pouch shipper", "RSC", 3, "B", 300, 200, 150, None, 400000),
    ("L07", "Toothpaste outer", "RSC", 3, "B", 360, 240, 160, None, 350000),
    ("L08", "Ketchup bottle outer", "RSC", 3, "C", 330, 250, 230, None, 250000),
    ("L09", "Baby food outer", "RSC", 3, "B", 340, 260, 200, None, 150000),
    ("L10", "E-commerce mailer", "RSC", 3, "B", 300, 250, 100, None, 300000),
    ("L11", "Shampoo sachet outer", "RSC", 3, "B", 400, 300, 200, None, 200000),
    ("L12", "Chocolate shipper", "RSC", 3, "B", 380, 280, 180, None, 220000),
    ("L13", "Detergent 1kg x12 outer", "RSC", 5, "BC", 420, 300, 280, None, 220000),
    ("L14", "Edible oil 1L x12 outer", "RSC", 5, "BC", 400, 300, 300, None, 300000),
    ("L15", "Juice 1L carton x12 outer", "RSC", 5, "BC", 380, 290, 260, None, 250000),
    ("L16", "Atta 5kg x4 outer", "RSC", 5, "BC", 450, 350, 300, None, 180000),
    ("L17", "Jam glass jar x24 outer", "RSC", 5, "BC", 420, 320, 200, None, 120000),
    ("L18", "PET 600ml x24 outer", "RSC", 5, "BC", 410, 280, 250, None, 260000),
    ("L19", "Rice 10kg x2 outer", "RSC", 5, "BC", 500, 350, 300, None, 90000),
    ("L20", "Floor cleaner 1L x12 outer", "RSC", 5, "BC", 400, 310, 290, None, 140000),
    ("L21", "Hair oil x24 outer", "RSC", 5, "BC", 380, 260, 220, None, 160000),
    ("L22", "Ghee 1L x12 outer", "RSC", 5, "BC", 400, 300, 260, None, 110000),
    ("L23", "Sauce bulk pack outer", "RSC", 5, "BC", 450, 320, 300, None, 80000),
    ("L24", "Pickle jar x12 outer", "RSC", 5, "BC", 360, 270, 220, None, 100000),
    ("L25", "Oil tin 15L outer", "RSC", 7, "BCB", 330, 330, 450, None, 60000),
    ("L26", "Bulk export master carton", "RSC", 7, "BCB", 600, 400, 400, None, 40000),
    ("L27", "24-cell partition set (for L17)", "PARTITION", 3, "B", 0, 0, 0, 0.40, 120000),
    ("L28", "24-cell partition set (for L18)", "PARTITION", 3, "B", 0, 0, 0, 0.42, 150000),
    ("L29", "Pallet layer pad 1200x1000", "PAD", 3, "C", 1200, 1000, 0, None, 50000),
    ("L30", "Divider pad 400x300", "PAD", 3, "B", 400, 300, 0, None, 400000),
]
MARKET_PER_KG = {3: 38.0, 5: 41.5, 7: 44.0}
LY_PER_KG = {3: 37.0, 5: 40.0, 7: 43.0}       # incumbent's 2025 contract basis
FX_USD_INR = 88.0                             # demo assumption, editable in the app
GST = 0.05                                    # [VERIFY] current GST rate on HSN 4819 boxes

lines = []
for (lid, desc, kind, ply, flute, L, W, H, area, qty) in LINES:
    if kind == "RSC":
        a = rsc_blank_m2(L, W, H)
    elif kind == "PAD":
        a = L * W / 1e6
    else:
        a = area
    gsm = board_gsm(ply, flute)
    wt = round(a * gsm / 1000, 4)
    if kind == "RSC":
        size = f"{L} x {W} x {H} mm (internal)"
    elif kind == "PAD":
        size = f"{L} x {W} mm"
    else:
        size = "24 cells, to suit outer"
    lines.append(dict(id=lid, description=desc, kind=kind, ply=ply, flute=flute, size=size,
                      L=L, W=W, H=H, board_gsm=gsm, weight_kg=wt, base_qty=qty,
                      grade=f"{ply}-ply {flute}-flute, layers {'/'.join(map(str, GRADES[ply]))} GSM, 18 BF",
                      print="1-colour flexo" if kind == "RSC" and ply < 7 else "plain"))

# scale annual quantities so the RFx is worth ~Rs 4 crore at market rates
mkt_value = sum(l["weight_kg"] * MARKET_PER_KG[l["ply"]] * l["base_qty"] for l in lines)
scale = 4.0e7 / mkt_value
for l in lines:
    l["annual_qty"] = int(round(l["base_qty"] * scale / 5000) * 5000)
    l["market_price"] = round(l["weight_kg"] * MARKET_PER_KG[l["ply"]], 2)
    l["last_year_price"] = round(l["weight_kg"] * LY_PER_KG[l["ply"]], 2)
    del l["base_qty"]

# ---------------------------------------------------------------- vendors
VENDORS = {
    "A": dict(name="Hosur Corrutech Pvt Ltd", city="Hosur", factor=0.978, noise=0.04),
    "B": dict(name="Pacific Fibre Packaging India Pvt Ltd", city="Nelamangala", factor=1.0, noise=0.03),
    "C": dict(name="Kaveri Kraft Boxes", city="Tumakuru", factor=0.94, noise=0.03),
    "D": dict(name="Sri Lakshmi Carton Works", city="Peenya, Bengaluru", factor=0.972, noise=0.05),
    "E": dict(name="Deccan Corrugators Pvt Ltd", city="Bidadi", factor=None, noise=0),
}
C_MISSING = {"L09", "L23", "L29"}
D_MISSING = {"L25", "L26"}
D_CORRECTED = "L18"          # printed price struck through, new price handwritten
C_ALTERNATE = "L14"          # also offers a lower-GSM alternate
E_PER_KG = {3: 38.0, 5: 42.0}  # straight from the brief's email

truth = {v: {} for v in VENDORS}   # normalised Rs/pc ex-GST, plus how it was stated
for l in lines:
    for v, meta in VENDORS.items():
        if v == "E":
            continue
        if (v == "C" and l["id"] in C_MISSING) or (v == "D" and l["id"] in D_MISSING):
            continue
        p = l["market_price"] * meta["factor"] * (1 + random.uniform(-meta["noise"], meta["noise"]))
        truth[v][l["id"]] = round(p, 2)

key = {"fx_usd_inr": FX_USD_INR, "gst": GST, "vendors": {}}


def rec(v, lid, stated, unit, currency, gst_incl, norm, flags=()):
    key["vendors"].setdefault(v, {"lines": {}})["lines"][lid] = dict(
        stated_price=stated, stated_unit=unit, currency=currency, gst_inclusive=gst_incl,
        normalised_inr_per_pc_ex_gst=round(norm, 2), expected_flags=list(flags))


# ---------- Vendor A: Excel that ignores the template, prices per 100 pcs, own order + wording
wb = Workbook(); ws = wb.active; ws.title = "Quotation"
ws.merge_cells("A1:G1"); ws["A1"] = "HOSUR CORRUTECH PVT LTD  -  QUOTATION"; ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A2:G2"); ws["A2"] = "SIPCOT Phase II, Hosur 635109  |  GSTIN 33AAHCH4821K1Z6  |  Ref: HCT/Q/26-27/0412  |  Dt 03.10.2026"
ws.merge_cells("A3:G3"); ws["A3"] = "Kind Attn: Ms. Ananya Rao, Category Buyer - Packaging   |   Your RFQ No. RFQ-PKG-2026-014"
hdr = ["Sl", "Item / Description", "Size (mm) LxWxH", "Ply / Flute", "Rate per 100 Nos (Rs.)", "GST %", "Remarks"]
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=5, column=i, value=h); c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(wrap_text=True)
order = sorted(lines, key=lambda l: (-l["ply"], l["id"]))  # they group heavy boxes first
r = 6
for i, l in enumerate(order, 1):
    p100 = round(truth["A"][l["id"]] * 100, 0)
    if l["kind"] == "RSC":
        d = f"{l['ply']} Ply Box ({l['flute']} flute)"; sz = f"{l['L']}x{l['W']}x{l['H']}"
    elif l["kind"] == "PAD":
        d = f"{l['ply']} Ply Pad"; sz = f"{l['L']}x{l['W']}"
    else:
        d = f"{l['ply']} Ply Partition 24 cell"; sz = "as per sample"
    remark = "Printing 1 col incl." if l["print"] != "plain" else ""
    if l["id"] == "L27": remark = "for jam jar box"
    if l["id"] == "L28": remark = "for PET bottle box"
    for c, val in enumerate([i, d, sz, f"{l['ply']}/{l['flute']}", f"{p100:,.2f}", "5%", remark], 1):
        ws.cell(row=r, column=c, value=val)
    rec("A", l["id"], p100, "per 100 pcs", "INR", False, p100 / 100, ["unit_converted"])
    r += 1
ws.cell(row=r + 1, column=2, value="Terms: Delivered Bidadi. GST extra @5%. Payment 60 days.")
ws.cell(row=r + 2, column=2, value="Rates valid 15 days from date of quote and subject to revision with kraft paper price movement.")
for col, w in zip("ABCDEFG", [5, 34, 18, 12, 20, 8, 24]):
    ws.column_dimensions[col].width = w
ws2 = wb.create_sheet("Tech Compliance")
qa_A = [("ISO 9001:2015", "Yes - certificate attached (valid till Aug 2027)"),
        ("BCT / Bursting strength reports", "Attached for 3, 5 and 7 ply grades"),
        ("Capacity", "300 MT per month, 2 corrugation lines"),
        ("Lead time", "7 days from PO"),
        ("Plant / distance", "Hosur, approx 75 km to Bidadi"),
        ("Kraft source", "Virgin + recycled mix, mills: BILT, Emami"),
        ("Food contact liner", "Available on request"),
        ("GST / MSME", "33AAHCH4821K1Z6 / Not MSME"),
        ("Payment terms", "60 days accepted"),
        ("References", "Can share 3 FMCG references on call")]
ws2.append(["Parameter", "Our response"])
for q in qa_A: ws2.append(list(q))
ws2.column_dimensions["A"].width = 32; ws2.column_dimensions["B"].width = 60
wb.save(VR / "A_HosurCorrutech_Quotation.xlsx")
key["vendors"]["A"]["format"] = "xlsx"
key["vendors"]["A"]["terms"] = dict(delivery="delivered Bidadi", gst="extra 5%", payment_days=60,
                                    validity="15 days, subject to kraft price revision")
key["vendors"]["A"]["vendor_flags"] = ["price_escalation_clause", "short_validity"]


# ---------- Vendor B: letterhead PDF, USD, footnote discount with a condition, total that doesn't add up
def pdf(path, story):
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                      topMargin=14 * mm, bottomMargin=14 * mm).build(story)


ss = getSampleStyleSheet()
st = ParagraphStyle("b", parent=ss["Normal"], fontName="Arial", fontSize=9, leading=12)
stb = ParagraphStyle("bb", parent=st, fontName="Arial-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#0B3D2E"))
sts = ParagraphStyle("s", parent=st, fontSize=6.5, leading=8, textColor=colors.HexColor("#666666"))
story = [Paragraph("PACIFIC FIBRE PACKAGING INDIA PVT LTD", stb),
         Paragraph("A Pacific Fibre Group company | Plot 22, KIADB Industrial Area, Nelamangala, Bengaluru Rural 562123", st),
         Spacer(1, 8),
         Paragraph("<b>Commercial Offer PFP/BLR/2026/1187</b> &nbsp;&nbsp; Date: 05 October 2026", st),
         Paragraph("To: Halasuru Consumer Products Pvt Ltd, Bidadi Plant. Against RFQ-PKG-2026-014.", st),
         Spacer(1, 6),
         Paragraph("We are pleased to offer as below. Group pricing policy requires quotation in US Dollars; invoicing in INR at RBI reference rate on invoice date.", st),
         Spacer(1, 6)]
rows = [["#", "Description", "Spec", "USD / pc", "Annual qty", "Indicative value USD"]]
tot = 0
for l in lines:
    usd = round(truth["B"][l["id"]] / FX_USD_INR, 3)
    val = round(usd * l["annual_qty"], 0)
    tot += val
    rows.append([l["id"][1:], l["description"], f"{l['ply']}P {l['flute']} {l['size'].split(' (')[0]}", f"{usd:.3f}",
                 f"{l['annual_qty']:,}", f"{val:,.0f}"])
    rec("B", l["id"], usd, "per pc", "USD", False, usd * FX_USD_INR, ["currency_converted"])
stated_total = tot + 1180  # arithmetic slip in their spreadsheet
rows.append(["", "TOTAL (indicative)", "", "", "", f"{stated_total:,.0f}*"])
t = Table(rows, colWidths=[8 * mm, 52 * mm, 50 * mm, 18 * mm, 22 * mm, 28 * mm], repeatRows=1)
t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Arial", 7.5), ("FONT", (0, 0), (-1, 0), "Arial-Bold", 7.5),
                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B3D2E")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                       ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("ALIGN", (3, 1), (-1, -1), "RIGHT")]))
story += [t, Spacer(1, 8),
          Paragraph("Terms: Prices FOR Bidadi, exclusive of GST. Payment: 45 days from invoice. Lead time: 10 working days. "
                    "Validity: 30 days. Quality: ISO 9001:2015 certified plant; certificate and test report enclosed.", st),
          Spacer(1, 14),
          Paragraph("For Pacific Fibre Packaging India Pvt Ltd &nbsp;&nbsp; Rohan Mehta, Key Account Manager", st),
          Spacer(1, 40),
          Paragraph("* Indicative total for reference only. A special volume discount of 2.5% on all line prices applies where the "
                    "annual purchase order value placed with Pacific Fibre exceeds INR 1.5 crore and payment is released within 45 days. "
                    "Not applicable to partial or trial orders. E&amp;OE.", sts)]
qa_B = [["Questionnaire", "Response"],
        ["ISO 9001:2015", "Yes, certified (see enclosure)"], ["Test reports", "BCT and burst reports enclosed"],
        ["Monthly capacity", "900 MT"], ["Lead time", "10 working days"], ["Plant & distance", "Nelamangala, ~55 km"],
        ["Kraft source", "Virgin kraft, partly imported"], ["Food-contact liner", "Yes"],
        ["GST / MSME", "29AAKCP5512M1Z3 / Large enterprise"], ["Payment terms", "45 days only (group policy)"],
        ["References", "Three national FMCG brands; names under NDA"]]
qt = Table(qa_B, colWidths=[50 * mm, 120 * mm])
qt.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Arial", 8), ("FONT", (0, 0), (-1, 0), "Arial-Bold", 8),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]))
story += [Spacer(1, 10), Paragraph("<b>Annexure: Supplier questionnaire</b>", st), qt]
pdf(VR / "B_PacificFibre_Offer.pdf", story)
key["vendors"]["B"].update(format="pdf", stated_total_usd=stated_total, true_total_usd=tot,
                           terms=dict(delivery="FOR Bidadi", gst="extra", payment_days=45, validity="30 days"),
                           conditional_discount=dict(pct=2.5, condition="annual PO value with vendor > INR 1.5 crore AND payment within 45 days"),
                           vendor_flags=["usd_quote", "footnote_discount_conditional", "total_does_not_reconcile",
                                         "iso_certificate_expired", "payment_terms_deviation_45_vs_60"])

# ---------- Vendor C: Word doc, commercials in prose, 27 of 30 lines, alternate spec, weak questionnaire
doc = Document()
s = doc.styles["Normal"]; s.font.name = "Calibri"; s.font.size = Pt(10.5)
doc.add_paragraph("KAVERI KRAFT BOXES\nB-14, KIADB Antharasanahalli, Tumakuru 572106 | Ph 0816-2291447")
doc.add_paragraph("Date: 06/10/2026\n\nTo,\nMs. Ananya Rao\nHalasuru Consumer Products Pvt Ltd, Bidadi\n\nSub: Our best rates against your RFQ-PKG-2026-014")
doc.add_paragraph("Dear Madam,\n\nThank you for considering Kaveri Kraft Boxes. We have studied your requirement carefully and are happy to quote our most competitive rates, which we are confident are the best in the market.")
groups = [("For the 3-ply shipper cartons", [l for l in lines if l["ply"] == 3 and l["kind"] == "RSC"]),
          ("For the 5-ply outer cartons in BC flute", [l for l in lines if l["ply"] == 5]),
          ("For the heavy 7-ply boxes", [l for l in lines if l["ply"] == 7]),
          ("For partitions and pads", [l for l in lines if l["kind"] in ("PARTITION", "PAD")])]
for intro, grp in groups:
    parts = []
    for l in grp:
        if l["id"] in C_MISSING:
            continue
        p = truth["C"][l["id"]]
        nm = l["size"].split(" (")[0] if l["kind"] != "PARTITION" else ("jam jar partition" if l["id"] == "L27" else "bottle partition")
        parts.append(f"{nm} at Rs.{p:.2f}")
        rec("C", l["id"], p, "per box", "INR", False, p)
    doc.add_paragraph(f"{intro}, our rate per box works out to " + ", ".join(parts) + ".")
alt = truth["C"][C_ALTERNATE]; alt_price = round(alt * 0.93, 2)
doc.add_paragraph(f"Further, for the edible oil outer (400x300x300) we can also offer an economical option with 120 GSM top liner "
                  f"instead of 180 GSM at only Rs.{alt_price:.2f}, which many customers are happily using.")
doc.add_paragraph("All above rates are plus GST 5%, delivered to your Bidadi plant. Payment 30 days preferred, 60 days can be considered for regular orders. "
                  "Delivery 12 days from PO. Rates firm for 3 months.")
doc.add_paragraph("Regarding your questionnaire: ISO certification has been applied for and the audit is scheduled for December 2026. "
                  "Test reports can be arranged on request. Our present capacity is 35 tonnes a month and we are expanding. "
                  "Kraft is sourced from local mills. We are a registered MSME (UDYAM-KR-27-0011842). References can be given on call.")
doc.add_paragraph("We request you to kindly give us an opportunity to serve you.\n\nThanking you,\nFor Kaveri Kraft Boxes\nS. Manjunath, Proprietor")
doc.save(VR / "C_KaveriKraft_Quotation.docx")
key["vendors"]["C"].update(format="docx", lines_quoted=27, missing=sorted(C_MISSING),
                           alternate_offer=dict(line=C_ALTERNATE, price=alt_price, deviation="120 GSM top liner vs 180 specified"),
                           terms=dict(delivery="delivered Bidadi", gst="extra 5%", payment_days="30 preferred, 60 negotiable", validity="3 months"),
                           vendor_flags=["missing_lines", "spec_deviation_alternate", "questionnaire_fail_iso", "questionnaire_fail_tests",
                                         "questionnaire_fail_capacity", "questionnaire_fail_lead_time"])

# ---------- Vendor D: printed rate card (incl. GST) -> photographed at an angle, handwritten correction
W_, H_ = 1700, 2400
card = Image.new("RGB", (W_, H_), "white"); dr = ImageDraw.Draw(card)
fB = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 54); fM = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 30)
fS = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 26); fH = ImageFont.truetype("C:/Windows/Fonts/Inkfree.ttf", 44)
dr.text((80, 70), "SRI LAKSHMI CARTON WORKS", font=fB, fill="black")
dr.text((80, 145), "No. 41, 3rd Stage, Peenya Industrial Area, Bengaluru 560058  |  Ph 98450 31277", font=fS, fill="black")
dr.text((80, 200), "RATE CARD - CORRUGATED BOXES   (Oct 2026)   For: Halasuru Consumer Products", font=fM, fill="black")
y = 270; dr.line((80, y, W_ - 80, y), fill="black", width=3); y += 15
dr.text((90, y), "Item", font=fM, fill="black"); dr.text((760, y), "Size / Ply", font=fM, fill="black"); dr.text((1330, y), "Rate Rs./pc", font=fM, fill="black")
y += 50; dr.line((80, y, W_ - 80, y), fill="black", width=2); y += 10
corr_xy = None
for l in lines:
    if l["id"] in D_MISSING:
        continue
    gross = round(truth["D"][l["id"]] * (1 + GST), 2)
    nm = l["description"][:30]
    spec = (l["size"].split(" (")[0].replace(" mm", "") if l["kind"] != "PARTITION" else "24 cell") + f"  {l['ply']}P"
    dr.text((90, y), nm, font=fS, fill="black"); dr.text((760, y), spec, font=fS, fill="black")
    if l["id"] == D_CORRECTED:
        old = round(gross * 0.94, 2)
        dr.text((1340, y), f"{old:,.2f}", font=fS, fill="black")
        corr_xy = (1340, y)
        rec("D", l["id"], gross, "per pc", "INR", True, truth["D"][l["id"]], ["handwritten_correction", "gst_inclusive"])
        key["vendors"]["D"]["handwritten"] = dict(line=l["id"], printed=old, corrected_to=gross)
    else:
        dr.text((1340, y), f"{gross:,.2f}", font=fS, fill="black")
        rec("D", l["id"], gross, "per pc", "INR", True, truth["D"][l["id"]], ["gst_inclusive", "from_photo"])
    y += 62
y += 20; dr.line((80, y, W_ - 80, y), fill="black", width=2); y += 20
for t_ in ["* All rates INCLUSIVE of GST @5%. Door delivery Bidadi free.", "* 7 ply not manufactured by us.",
           "* Min. order 10,000 pcs per size.  Monthly capacity 60 tonnes.", "* Payment 30 days.  Validity 1 month."]:
    dr.text((90, y), t_, font=fS, fill="black"); y += 42
card.save(VR / "D_rate_card_PRINT_THIS.png")
# handwritten correction drawn on a copy for the synthetic photo (when you print, write it by pen instead)
hand = card.copy(); hd = ImageDraw.Draw(hand)
x0, y0 = corr_xy
hd.line((x0 - 5, y0 + 16, x0 + 120, y0 + 12), fill=(20, 40, 140), width=5)
hd.text((x0 + 140, y0 - 14), f"{key['vendors']['D']['handwritten']['corrected_to']:.2f}", font=fH, fill=(20, 40, 140))
# perspective warp + shading + blur = phone photo at an angle
bg = Image.new("RGB", (2200, 2800), (92, 78, 64))
src = hand.resize((1700, 2400))
coeffs_target = [(260, 180), (1930, 330), (2050, 2620), (120, 2480)]


def find_coeffs(pa, pb):
    import numpy as np
    m = []
    for p1, p2 in zip(pa, pb):
        m.append([p1[0], p1[1], 1, 0, 0, 0, -p2[0] * p1[0], -p2[0] * p1[1]])
        m.append([0, 0, 0, p1[0], p1[1], 1, -p2[1] * p1[0], -p2[1] * p1[1]])
    A = np.array(m, dtype=float); B = np.array(pb, dtype=float).reshape(8)
    return list(np.linalg.solve(A, B))


co = find_coeffs(coeffs_target, [(0, 0), (1700, 0), (1700, 2400), (0, 2400)])
warped = src.transform((2200, 2800), Image.PERSPECTIVE, co, Image.BICUBIC, fillcolor=(92, 78, 64))
mask = Image.new("L", (1700, 2400), 255).transform((2200, 2800), Image.PERSPECTIVE, co, Image.BICUBIC, fillcolor=0)
bg.paste(warped, (0, 0), mask)
shade = Image.linear_gradient("L").resize((2200, 2800)).point(lambda v: int(v * 0.35))
bg = Image.composite(Image.new("RGB", bg.size, (0, 0, 0)), bg, shade)
bg = bg.filter(ImageFilter.GaussianBlur(1.6)).resize((1100, 1400))
bg.save(VR / "D_SriLakshmi_RateCard_photo.jpg", quality=72)
key["vendors"]["D"].update(format="jpg photo", lines_quoted=28, missing=sorted(D_MISSING),
                           terms=dict(delivery="free door delivery Bidadi", gst="INCLUSIVE 5%", payment_days=30, validity="1 month"),
                           vendor_flags=["gst_inclusive", "photo_angled", "handwritten_correction", "no_7ply", "capacity_60t"])
(VR / "D_SriLakshmi_email.txt").write_text(
    "From: Sri Lakshmi Carton Works <enquiry@srilakshmicartons.example>\nTo: ananya.rao@halasuru-cp.example\nDate: Tue, 7 Oct 2026 19:42\n"
    "Subject: RE: RFQ-PKG-2026-014 rates\n\nMadam,\nPFA our rate card photo. Typed sheet will take time, printer not working.\n"
    "Questionnaire: ISO 9001 yes cert attached. BCT test report attached. Capacity 60 ton/month. Lead time 7 days. "
    "Peenya to Bidadi 40 km. Kraft from local mills. MSME yes. Payment 30 days pls. Customers - 2 local FMCG + 1 pharma.\n\n"
    "Regards\nRaghavendra K\n", encoding="utf-8")

# ---------- Vendor E: incumbent, one-line email
ly = {l["id"]: l["last_year_price"] for l in lines}
(VR / "E_DeccanCorrugators_email.txt").write_text(
    "From: Venkatesh Gowda <venkatesh@deccancorrugators.example>\nTo: ananya.rao@halasuru-cp.example\n"
    "Date: Fri, 10 Oct 2026 21:07\nSubject: Re: RFQ-PKG-2026-014\n\n"
    "Ananya ji,\n\nRs.42/kg for the 5-ply, 38 for the 3-ply, rest same as last year, freight extra.\n\n"
    "Venkatesh\nDeccan Corrugators\nSent from my iPhone\n", encoding="utf-8")
for l in lines:
    if l["kind"] == "RSC" and l["ply"] in E_PER_KG:
        rk = E_PER_KG[l["ply"]]
        rec("E", l["id"], rk, "per kg", "INR", False, rk * l["weight_kg"], ["per_kg_derived", "freight_extra"])
    elif l["kind"] == "RSC":  # 7-ply: genuinely "rest"
        rec("E", l["id"], ly[l["id"]], "same as last year", "INR", False, ly[l["id"]], ["inferred_from_last_year", "freight_extra"])
    else:  # partitions and pads are 3-ply board: 38/kg OR last year? genuinely ambiguous
        rec("E", l["id"], ly[l["id"]], "AMBIGUOUS: 3-ply per-kg rate or last year", "INR", False, ly[l["id"]],
            ["ambiguous_scope", "inferred_from_last_year", "freight_extra"])
        key["vendors"]["E"]["lines"][l["id"]]["alt_if_per_kg"] = round(38.0 * l["weight_kg"], 2)
key["vendors"]["E"].update(format="email text", incumbent=True,
                           terms=dict(delivery="freight EXTRA (last year: delivered)", gst="not stated, assume extra",
                                      payment_days="not stated (last year 60)", validity="not stated"),
                           vendor_flags=["per_kg_pricing", "rest_same_as_last_year", "freight_extra", "questionnaire_not_answered",
                                         "partition_pad_scope_ambiguous"])

# ---------------------------------------------------------------- attachments


def cert(path, title, body_lines):
    pdf(path, [Paragraph(title, stb), Spacer(1, 10)] + [Paragraph(b, st) for b in body_lines])


cert(AT / "A_ISO9001_certificate.pdf", "CERTIFICATE OF REGISTRATION - ISO 9001:2015",
     ["Issued to: Hosur Corrutech Pvt Ltd, SIPCOT Phase II, Hosur", "Scope: Manufacture of corrugated boxes, sheets and partitions",
      "Certificate No: QMS/IN/88231", "Issue date: 14 August 2024", "<b>Valid until: 13 August 2027</b>", "Certification body: Sapient Assurance Registrar (fictional)"])
cert(AT / "A_BCT_test_report.pdf", "TEST REPORT - BOX COMPRESSION & BURST",
     ["Supplier: Hosur Corrutech Pvt Ltd", "Report date: 22 September 2026",
      "3-ply 18BF: Burst 7.4 kg/cm2, BCT 290 kgf (400x280x220)", "5-ply 18BF: Burst 12.8 kg/cm2, BCT 610 kgf (400x300x300)",
      "7-ply 18BF: Burst 17.1 kg/cm2, BCT 980 kgf (600x400x400)", "Result: PASS against customer spec"])
cert(AT / "B_ISO9001_certificate.pdf", "CERTIFICATE OF REGISTRATION - ISO 9001:2015",
     ["Issued to: Pacific Fibre Packaging India Pvt Ltd, Nelamangala", "Scope: Corrugated packaging", "Certificate No: IN-QM-55102",
      "Issue date: 1 April 2023", "<b>Valid until: 31 March 2026</b>", "Certification body: Meridian Quality Certification (fictional)"])
cert(AT / "B_test_report.pdf", "QUALITY TEST SUMMARY", ["Supplier: Pacific Fibre Packaging India Pvt Ltd", "Report date: 30 June 2026",
                                                        "Burst and BCT within spec for 3/5/7 ply grades. Moisture 7.2%.", "Result: PASS"])
cert(AT / "D_ISO9001_certificate.pdf", "CERTIFICATE OF REGISTRATION - ISO 9001:2015",
     ["Issued to: Sri Lakshmi Carton Works, Peenya, Bengaluru", "Certificate No: QA-KA-30917", "Issue date: 2 January 2026",
      "<b>Valid until: 1 January 2029</b>", "Certification body: Sapient Assurance Registrar (fictional)"])
cert(AT / "D_BCT_test_report.pdf", "BURST / BCT TEST REPORT", ["Supplier: Sri Lakshmi Carton Works", "Report date: 1 October 2026",
                                                               "3-ply: Burst 7.1 kg/cm2. 5-ply: Burst 12.2 kg/cm2.", "Result: PASS"])
key["attachments"] = {"A": ["A_ISO9001_certificate.pdf", "A_BCT_test_report.pdf"],
                      "B": ["B_ISO9001_certificate.pdf (EXPIRED 2026-03-31)", "B_test_report.pdf"],
                      "C": [], "D": ["D_ISO9001_certificate.pdf", "D_BCT_test_report.pdf"], "E": []}
key["questionnaire_outcome"] = {"A": "pass", "B": "issue: ISO certificate expired 2026-03-31; payment 45 days vs 60",
                                "C": "fail: no ISO, no test reports, capacity 35t < 50t, lead time 12d > 10d",
                                "D": "pass (note: capacity 60t/month limits share of award)", "E": "unknown: not answered"}

# ---------------------------------------------------------------- RFx + last year
QUESTIONNAIRE = [
    dict(id="Q1", text="Is your plant ISO 9001:2015 certified? Attach the current certificate.", must_pass=True, rule="valid certificate on award date"),
    dict(id="Q2", text="Attach BCT and bursting strength test reports for the offered grades.", must_pass=True, rule="reports provided"),
    dict(id="Q3", text="Monthly corrugation capacity (tonnes).", must_pass=True, rule=">= 50 tonnes/month"),
    dict(id="Q4", text="Standard lead time from PO (days).", must_pass=True, rule="<= 10 days"),
    dict(id="Q5", text="Plant location and distance to Bidadi plant.", must_pass=False, rule=None),
    dict(id="Q6", text="Kraft paper source (virgin / recycled, mills).", must_pass=False, rule=None),
    dict(id="Q7", text="Food-contact compliant inner liner available?", must_pass=False, rule=None),
    dict(id="Q8", text="GSTIN and MSME status.", must_pass=False, rule=None),
    dict(id="Q9", text="Do you accept 60-day payment terms?", must_pass=False, rule="preferred"),
    dict(id="Q10", text="Three current FMCG customer references.", must_pass=False, rule=None),
]
rfx = dict(
    id="RFQ-PKG-2026-014", title="Annual rate contract - corrugated boxes, partitions and pads (FY 2026-27)",
    buyer=dict(company="Halasuru Consumer Products Pvt Ltd (fictional)", person="Ananya Rao, Category Buyer - Packaging",
               approver="VP Procurement", plant="Bidadi Industrial Area, Bengaluru"),
    issued="2026-10-01", due="2026-10-10",
    terms=dict(currency="INR", prices="ex-GST, per piece", delivery="FOR Bidadi plant (freight included)",
               payment_days=60, validity="12 months", price_basis="firm for 6 months, then review against kraft index"),
    lines=lines, questionnaire=QUESTIONNAIRE,
    vendors={k: dict(name=v["name"], city=v["city"]) for k, v in VENDORS.items()},
    incumbent="E",
    assumptions=dict(fx_usd_inr=FX_USD_INR, gst_pct=GST * 100, freight_pct_if_extra=3.0,
                     board_formula="RSC blank = (2L+2W+35) x (W+H) mm; weight = area x board GSM; flute take-up B 1.35, C 1.45"),
)
(ROOT / "rfx.json").write_text(json.dumps(rfx, indent=2), encoding="utf-8")
with open(ROOT / "last_year_contract.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["line_id", "description", "vendor", "price_inr_per_pc_ex_gst", "annual_qty_2025", "contract", "freight"])
    for l in lines:
        w.writerow([l["id"], l["description"], "Deccan Corrugators Pvt Ltd", l["last_year_price"],
                    int(l["annual_qty"] * 0.92 // 5000 * 5000), "RC-2025-PKG-007", "delivered (included)"])
(ROOT / "answer_key.json").write_text(json.dumps(key, indent=2), encoding="utf-8")

total = sum(l["market_price"] * l["annual_qty"] for l in lines)
print(f"lines={len(lines)} market value=Rs {total/1e7:.2f} cr, tonnes/yr={sum(l['weight_kg']*l['annual_qty'] for l in lines)/1000:.0f}")
for v in "ABCDE":
    ls_ = key["vendors"][v]["lines"]
    print(v, len(ls_), "lines, value Rs", round(sum(ls_[i]["normalised_inr_per_pc_ex_gst"] * next(l["annual_qty"] for l in lines if l["id"] == i) for i in ls_) / 1e7, 2), "cr")
